"""`document.extract` — the text of a PDF, DOCX, ODT, image (OCR) or plain-text file.

A scanned PDF has no text layer, so `get_text()` returns "" for every page and the document
would be indexed as empty — no error, simply never findable. `ocr: auto` rasterizes any page
that yields (almost) nothing and runs Tesseract on it; `always` OCRs every page; `never` reads
only the text layer.

OCR on a photo answers with pages of plausible-looking garbage rather than with nothing, so an
image's OCR keeps only confident multi-character words and is reported as empty text below a
floor of alphanumerics. Both conditions are needed: noise scores very high on single glyphs.
"""

from __future__ import annotations

import asyncio
import os
import tempfile
from typing import Literal

from fastapi import APIRouter, Request

from ..context import Context
from ..errors import McsError, TOOL_FAILED, UNSUPPORTED
from ..schemas import FileRef, WithOptions
from ..streaming import Outcome, Progress, run_op
from ..tools.decode import is_raw
from ..tools.run import require, run
from ..tools.versions import tool_version
from .media import kind_of

# Tesseract parallelises one page across cores with OpenMP (four threads by default), for a
# gain of a few percent; in the tool lane that is four threads per slot on a box that is also
# decoding, encoding and running models. One thread per OCR, like every other tool here. Set
# on the front's own environment, which `tools/run.py` hands to every tool it spawns; the model
# worker is spawned without it (worker.TOOL_ONLY_ENV), since torch on the CPU wants the cores.
os.environ.setdefault("OMP_THREAD_LIMIT", "1")

# tesseract used to be spawned by pytesseract from inside a thread. That cost it the tools
# wrapper's `nice -n 10` (a document crawl once ran four tesseracts at the front's priority on
# a six-core host — load average 22 — and the model worker could not get scheduled inside a
# search box's two-second budget), and it made the run uncancellable: a caller that gave up
# and closed its connection reached the op's task, not a subprocess owned by a thread, so the
# binary ran on as an orphan holding the tools slot. It goes through `tools/run.py` now like
# every other tool: niced, in its own process group, and killed on timeout or cancellation.

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
ODT = "application/vnd.oasis.opendocument.text"

OCR_DPI = int(os.environ.get("MCS_OCR_DPI", "200"))
MIN_PAGE_TEXT_CHARS = int(os.environ.get("MCS_OCR_MIN_PAGE_CHARS", "32"))
MAX_OCR_PAGES = int(os.environ.get("MCS_OCR_MAX_PAGES", "40"))
MIN_IMAGE_TEXT_CHARS = int(os.environ.get("MCS_OCR_MIN_IMAGE_CHARS", "12"))
MIN_WORD_CONF = float(os.environ.get("MCS_OCR_MIN_WORD_CONF", "60"))
# The longest one tesseract run may take. Photos of textured surfaces — a renovation's bare
# walls, five of them on fry2 on 2026-09-16 — send tesseract into a pass that has not finished
# after twenty-five minutes at 70 % of a core. The caller gives up at its own deadline, but a
# closed connection does not reach a binary pytesseract spawned inside a thread, so each such
# image left an orphan holding a tools slot: five images, two attempts, and the lane was full
# for three hours while fingerprints and renditions queued behind them into their own
# deadlines. Past this the run is killed and the image counts as having no text, which for a
# photo is the honest answer; a scanned page that genuinely needs longer is what the setting
# is for.
OCR_TIMEOUT_SECONDS = float(os.environ.get("MCS_OCR_TIMEOUT_SECONDS", "120"))

# What leptonica opens on its own. Anything else (HEIC, an odd TIFF) is decoded by PIL into a
# PNG for the run. MPO is a JPEG with extra frames; leptonica reads the first.
TESSERACT_READS = frozenset(["JPEG", "MPO", "PNG", "TIFF", "BMP", "GIF", "WEBP", "PNM", "PPM"])
MIN_WORD_CHARS = 2
DEFAULT_LANGUAGES = ("swe", "eng")

# Failures that are a property of the file, not of the run. Named by module.class so this
# costs no imports; deliberately narrow — anything unrecognised stays transient and loud.
PERMANENT_ERRORS = frozenset([
    "PIL.UnidentifiedImageError",
    "pymupdf.FileDataError",
    "pymupdf.EmptyFileError",
    "zipfile.BadZipFile",
])


class DocumentRequest(WithOptions):
    file: FileRef
    ocr: Literal["auto", "always", "never"] = "auto"
    languages: list[str] | None = None


def is_permanent(exc: Exception) -> bool:
    return f"{type(exc).__module__}.{type(exc).__name__}" in PERMANENT_ERRORS


async def tesseract(source: str, lang: str, *, tsv: bool) -> str | None:
    """One tesseract run over an image file: its stdout, or None when it ran out of time.

    A timeout is not an error here — see OCR_TIMEOUT_SECONDS — but a cancelled request still
    propagates, with the process group already killed by `run`.
    """
    args = [require("tesseract"), source, "stdout", "-l", lang, *(["tsv"] if tsv else [])]
    try:
        done = await run(args, timeout=OCR_TIMEOUT_SECONDS)
    except asyncio.TimeoutError:
        return None
    if done.code != 0:
        raise McsError(TOOL_FAILED, f"tesseract failed: {done.tail() or f'exit {done.code}'}")
    return done.stdout.decode(errors="replace")


def parse_tsv(text: str) -> dict[str, list]:
    """tesseract's `tsv` output as columns, the shape pytesseract's `image_to_data` gave."""
    lines = [line for line in text.splitlines() if line.strip()]
    if not lines:
        return {"text": [], "conf": [], "page_num": [], "block_num": [], "par_num": [], "line_num": []}
    header = lines[0].split("\t")
    columns: dict[str, list] = {name: [] for name in header}
    for line in lines[1:]:
        cells = line.split("\t")
        if len(cells) < len(header):
            cells += [""] * (len(header) - len(cells))
        for name, cell in zip(header, cells):
            if name == "text":
                columns[name].append(cell)
            elif name == "conf":
                columns[name].append(float(cell) if cell not in ("", "-") else -1.0)
            else:
                columns[name].append(int(cell) if cell.lstrip("-").isdigit() else 0)
    return columns


def _pixmap_png(page) -> bytes:
    import fitz

    return page.get_pixmap(matrix=fitz.Matrix(OCR_DPI / 72, OCR_DPI / 72)).tobytes("png")


async def _ocr_page(page, lang: str) -> str:
    try:
        png = await asyncio.to_thread(_pixmap_png, page)
    except Exception:  # noqa: BLE001 - OCR is a bonus, never fatal
        return ""
    with tempfile.NamedTemporaryFile(prefix="mcs-ocr-", suffix=".png", delete=False) as handle:
        handle.write(png)
        source = handle.name
    try:
        try:
            out = await tesseract(source, lang, tsv=False)
        except McsError:
            return ""
        return (out or "").strip()
    finally:
        try:
            os.unlink(source)
        except OSError:
            pass


async def extract_pdf(path: str, ocr: str, lang: str) -> tuple[str, int, str]:
    import fitz

    doc = await asyncio.to_thread(fitz.open, path)
    pages = len(doc)
    parts: list[str] = []
    ocr_pages = 0
    try:
        for page in doc:
            text = (await asyncio.to_thread(page.get_text)).strip() if ocr != "always" else ""
            wants_ocr = ocr == "always" or (ocr == "auto" and len(text) < MIN_PAGE_TEXT_CHARS)
            if wants_ocr and ocr_pages < MAX_OCR_PAGES:
                recognised = await _ocr_page(page, lang)
                if len(recognised) > len(text):
                    text = recognised
                    ocr_pages += 1
            if text:
                parts.append(text)
    finally:
        doc.close()
    return "\n\n".join(parts), pages, "ocr" if ocr_pages and ocr_pages * 2 >= max(1, pages) else "pdf-text"


def extract_docx(path: str) -> tuple[str, int, str]:
    from docx import Document

    doc = Document(path)
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts), 0, "docx"


def extract_odt(path: str) -> tuple[str, int, str]:
    from odf.opendocument import load
    from odf.text import P

    doc = load(path)
    parts = []
    for elem in doc.getElementsByType(P):
        text = "".join(getattr(node, "data", None) or str(node) for node in elem.childNodes)
        if text.strip():
            parts.append(text.strip())
    return "\n\n".join(parts), 0, "odf"


def _confident_text(data: dict) -> str:
    lines: dict[tuple, list[str]] = {}
    for index, word in enumerate(data["text"]):
        word = (word or "").strip()
        if not word or float(data["conf"][index]) < MIN_WORD_CONF:
            continue
        if sum(c.isalnum() for c in word) < MIN_WORD_CHARS:
            continue
        key = (data["page_num"][index], data["block_num"][index], data["par_num"][index], data["line_num"][index])
        lines.setdefault(key, []).append(word)
    return "\n".join(" ".join(words) for _, words in sorted(lines.items())).strip()


def _ocr_source(path: str) -> tuple[str, bool]:
    """The file tesseract should read for `path`: the file itself, or a PNG decoded by PIL.

    Returns (source, temporary). Opening with PIL first is what applies the format check and
    the bomb guard; the decode itself only happens for a format leptonica cannot read.
    """
    from PIL import Image, ImageFile

    ImageFile.LOAD_TRUNCATED_IMAGES = True
    # A high-dpi album scan trips PIL's decompression-bomb guard, which exists for untrusted
    # uploads this service never sees. Bounded rather than disabled so a corrupt header fails.
    Image.MAX_IMAGE_PIXELS = 500_000_000
    with Image.open(path) as image:
        if image.format in TESSERACT_READS:
            return path, False
        with tempfile.NamedTemporaryFile(prefix="mcs-ocr-", suffix=".png", delete=False) as handle:
            image.convert("RGB").save(handle, format="PNG")
            return handle.name, True


async def extract_image(path: str, lang: str) -> tuple[str, int, str]:
    source, temporary = await asyncio.to_thread(_ocr_source, path)
    try:
        out = await tesseract(source, lang, tsv=True)
    finally:
        if temporary:
            try:
                os.unlink(source)
            except OSError:
                pass
    if out is None:
        return "", 0, "ocr"
    text = _confident_text(parse_tsv(out))
    if sum(c.isalnum() for c in text) < MIN_IMAGE_TEXT_CHARS:
        return "", 0, "ocr"
    return text, 0, "ocr"


def extract_plaintext(path: str) -> tuple[str, int, str]:
    with open(path, encoding="utf-8", errors="replace") as handle:
        return handle.read().strip(), 0, "text"


async def extract(path: str, mimetype: str | None, ocr: str, lang: str) -> tuple[str, int, str]:
    if mimetype == "application/pdf":
        return await extract_pdf(path, ocr, lang)
    if mimetype == DOCX:
        return await asyncio.to_thread(extract_docx, path)
    if mimetype == ODT:
        return await asyncio.to_thread(extract_odt, path)
    if is_raw(mimetype):
        # A sensor dump has no text to read, and PIL cannot open one: the libtiff reader
        # fails with an OSError ("Error setting from dictionary") that is not in
        # PERMANENT_ERRORS, so 81 DNGs on fry2 each cost three attempts and a failed
        # instance. Refused up front, permanently, like any other format with no extractor.
        raise McsError(UNSUPPORTED, f"no text extractor for camera RAW ({mimetype})")
    if mimetype and mimetype.startswith("image/"):
        if ocr == "never":
            return "", 0, "ocr"
        return await extract_image(path, lang)
    if mimetype and mimetype.startswith("text/"):
        return await asyncio.to_thread(extract_plaintext, path)
    raise McsError(UNSUPPORTED, f"no text extractor for {mimetype or 'an unknown type'}")


async def run_extract(ctx: Context, req: DocumentRequest, progress: Progress) -> Outcome:
    path = ctx.roots.input(req.file.path)
    mimetype = req.file.mimetype
    lang = "+".join(req.languages or DEFAULT_LANGUAGES)
    async with ctx.admission.slot("tools", interactive=req.options.interactive):
        if not mimetype and ctx.exiftool is not None:
            try:
                mimetype = str((await ctx.exiftool.extract(path)).get("MIMEType") or "")
            except McsError:
                mimetype = None
        if not mimetype:
            import mimetypes

            mimetype = mimetypes.guess_type(path)[0]
        await progress.emit("extract", message=kind_of(mimetype))
        try:
            text, pages, method = await extract(path, mimetype, req.ocr, lang)
        except McsError:
            raise
        except Exception as exc:  # noqa: BLE001 - classified below
            raise McsError(TOOL_FAILED, str(exc) or type(exc).__name__, permanent=is_permanent(exc)) from exc
    result = {"text": text, "pages": pages, "method": method, "word_count": len(text.split())}
    tess = await tool_version("tesseract") if method == "ocr" else None
    return Outcome(result, producer=f"tesseract/{tess}" if tess else method)


def register(router: APIRouter, ctx: Context) -> None:
    @router.post("/document/extract")
    async def document_extract(request: Request, body: DocumentRequest):
        return await run_op(request, lambda progress: run_extract(ctx, body, progress))
